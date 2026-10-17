from pathlib import Path
import re, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).resolve().parent
sql = (root / 'questions.sql').read_text(encoding='utf-8-sig')
pattern = r"\((\d+),'((?:[^']|'')*)','((?:[^']|'')*)'::jsonb,(\d+)\)"
questions = [(int(i), text.replace("''", "'"), json.loads(options.replace("''", "'")), int(answer)) for i, text, options, answer in re.findall(pattern, sql)]
assert [q[0] for q in questions] == list(range(1,41))
assert all(len(q[2]) == 5 and 0 <= q[3] < 5 for q in questions)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.7)
section.left_margin, section.right_margin = Cm(2), Cm(2)
for name in ('Normal','Title','Heading 1','Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0
doc.styles['Title'].font.size = Pt(21)
doc.styles['Title'].paragraph_format.space_after = Pt(8)
doc.add_paragraph('AutoCAD Champions', 'Title')
p=doc.add_paragraph('Banco de 40 questões para revisão')
p.runs[0].bold=True
p.paragraph_format.space_after=Pt(6)
doc.add_paragraph('Enunciados, alternativas e gabaritos cadastrados no arquivo de perguntas do campeonato. Use a numeração original para indicar as questões que deseja substituir ou modificar.')
p=doc.add_paragraph('Revisão do professor  •  16 de outubro de 2026')
p.paragraph_format.space_after=Pt(13)
for index, (number, text, options, answer) in enumerate(questions):
    if index and index % 4 == 0:
        doc.add_page_break()
        p=doc.add_paragraph(f'AutoCAD Champions   |   Questões {number:02d} a {min(number+3,40):02d}')
        p.runs[0].bold=True
        p.paragraph_format.space_after=Pt(14)
    p=doc.add_paragraph(text)
    p.runs[0].bold=True
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_with_next=True
    for option in options:
        p=doc.add_paragraph(option)
        p.paragraph_format.left_indent=Cm(.25)
        p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph('Gabarito atual: '+options[answer])
    p.runs[0].bold=True
    p.paragraph_format.space_before=Pt(5)
    p.paragraph_format.keep_with_next=True
    p=doc.add_paragraph('Anotações: __________________________________________________________')
    p.paragraph_format.space_before=Pt(5)
    p.paragraph_format.space_after=Pt(14)
footer=section.footer.paragraphs[0]
footer.alignment=2
footer.add_run('AutoCAD Champions  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for run in footer.runs: run.font.size=Pt(8)
doc.core_properties.title='AutoCAD Champions banco de 40 questões para revisão'
doc.core_properties.subject='Questões, alternativas e gabaritos atuais para revisão do professor'
output=root/'AutoCAD_Champions_40_Questoes_Revisao.docx'
doc.save(output)
# Conferência do conteúdo após salvar o documento.
saved=Document(output)
paragraphs=[p.text for p in saved.paragraphs]
for _, text, options, answer in questions:
    assert text in paragraphs
    assert all(option in paragraphs for option in options)
    assert 'Gabarito atual: '+options[answer] in paragraphs
print('Documento criado e conferido: 40 questões, 200 alternativas e 40 gabaritos.')
